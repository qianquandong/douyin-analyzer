import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

from .utils import format_timestamp


def generate_report(
    video_info: dict,
    url: str,
    scenes: list[dict],
    transcription: dict | None,
    analyses: list[dict] | None,
    output_dir: str,
) -> str:
    """
    生成Markdown分析报告。

    返回报告文件路径。
    """
    report_path = os.path.join(output_dir, "report.md")
    lines = []

    # 标题
    lines.append("# 抖音视频分析报告\n")

    # 视频信息
    lines.append("## 视频信息\n")
    lines.append(f"- **标题**: {video_info.get('title', '未知')}")
    lines.append(f"- **链接**: {url}")
    duration = video_info.get("duration", 0)
    lines.append(f"- **时长**: {format_timestamp(duration)}")
    lines.append(f"- **场景数**: {len(scenes)}")
    lines.append(f"- **分析时间**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append("")

    # 时间线
    lines.append("## 场景时间线\n")

    for scene in scenes:
        num = scene["scene_num"]
        start = format_timestamp(scene["start"])
        end = format_timestamp(scene["end"])
        frame_paths = scene["frame_paths"]

        lines.append(f"### 场景 {num} ({start} - {end})\n")

        # 显示所有关键帧截图
        for j, fp in enumerate(frame_paths):
            rel_path = os.path.relpath(fp, output_dir)
            if len(frame_paths) == 1:
                lines.append(f"![场景{num}]({rel_path})\n")
            else:
                label = "首帧" if j == 0 else "尾帧"
                lines.append(f"**{label}**: ![场景{num}-{label}]({rel_path})\n")

        # 转录文本
        if transcription and transcription.get("segments"):
            seg_text = _get_segments_for_scene(
                transcription["segments"], scene["start"], scene["end"]
            )
            if seg_text:
                lines.append(f"**转录文本**: {seg_text}\n")
            else:
                lines.append("**转录文本**: （此段无语音）\n")

        # AI分析（每张帧图都有对应的分析）
        if analyses:
            for fp in frame_paths:
                analysis = _find_analysis_for_scene(analyses, fp)
                if analysis:
                    rel_path = os.path.relpath(fp, output_dir)
                    if len(frame_paths) > 1:
                        fname = os.path.basename(fp)
                        lines.append(f"#### {fname}\n")
                    lines.append(f"**画面描述词**: {analysis['description']}\n")
                    lines.append("**AI提示词（中文）**:")
                    lines.append(f"> {analysis['prompt']}\n")

        lines.append("---\n")

    # 完整转录
    if transcription and transcription.get("text"):
        lines.append("## 完整转录文本\n")
        lines.append(transcription["text"])
        lines.append("")

    with open(report_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    return report_path


def _get_segments_for_scene(
    segments: list[dict], start: float, end: float
) -> str:
    """获取属于某个场景时间范围内的转录文本。"""
    texts = []
    for seg in segments:
        seg_mid = (seg["start"] + seg["end"]) / 2
        if start <= seg_mid <= end:
            texts.append(seg["text"].strip())
    return "".join(texts)


def _find_analysis_for_scene(
    analyses: list[dict], frame_path: str
) -> dict | None:
    """根据帧路径查找对应的分析结果。"""
    for analysis in analyses:
        if analysis.get("frame_path") == frame_path:
            return analysis
    return None


def _make_thumbnail(src_path: str, output_dir: str, max_h: int = 120) -> str:
    """生成缩略图用于嵌入Excel。返回缩略图路径。"""
    thumb_dir = os.path.join(output_dir, "frames", ".thumbs")
    os.makedirs(thumb_dir, exist_ok=True)
    thumb_path = os.path.join(thumb_dir, os.path.basename(src_path))
    img = PILImage.open(src_path)
    ratio = max_h / img.height
    new_size = (int(img.width * ratio), max_h)
    img = img.resize(new_size, PILImage.LANCZOS)
    img.save(thumb_path, "JPEG", quality=80)
    return thumb_path


def generate_excel(
    video_info: dict,
    url: str,
    scenes: list[dict],
    transcription: dict | None,
    analyses: list[dict] | None,
    output_dir: str,
) -> str:
    """生成Excel分析报告（含嵌入图片）。返回文件路径。"""
    xlsx_path = os.path.join(output_dir, "report.xlsx")
    wb = Workbook()

    IMG_COL_1 = 6      # F列放第1帧
    IMG_COL_2 = 7      # G列放第2帧（双帧场景）
    IMG_HEIGHT = 120    # 缩略图高度(px)
    ROW_HEIGHT_PTS = 100  # 行高(pt)

    # ── Sheet 1: 视频分析 ──
    ws = wb.active
    ws.title = "视频分析"

    headers = [
        "场景编号", "开始时间", "结束时间", "时长(秒)",
        "帧数", "关键帧1", "关键帧2", "转录文本",
        "画面描述词(中文)", "AI提示词(中文)",
    ]
    h_fill = PatternFill("solid", fgColor="1F4E79")
    h_font = Font(bold=True, color="FFFFFF", name="Arial", size=11)
    h_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(
        left=Side("thin", color="D9D9D9"),
        right=Side("thin", color="D9D9D9"),
        top=Side("thin", color="D9D9D9"),
        bottom=Side("thin", color="D9D9D9"),
    )
    alt_fill = PatternFill("solid", fgColor="F2F7FB")
    body_font = Font(name="Arial", size=10)
    wrap = Alignment(vertical="top", wrap_text=True)
    center = Alignment(horizontal="center", vertical="center")

    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill, c.font, c.alignment, c.border = h_fill, h_font, h_align, border

    for i, scene in enumerate(scenes):
        row = i + 2
        fill = alt_fill if i % 2 == 0 else PatternFill()
        start_t = format_timestamp(scene["start"])
        end_t = format_timestamp(scene["end"])
        dur = round(scene["end"] - scene["start"], 1)

        # 转录文本
        trans = ""
        if transcription and transcription.get("segments"):
            trans = _get_segments_for_scene(
                transcription["segments"], scene["start"], scene["end"]
            )

        # AI分析
        descs, prompts = [], []
        if analyses:
            for fp in scene["frame_paths"]:
                a = _find_analysis_for_scene(analyses, fp)
                if a:
                    descs.append(a["description"])
                    prompts.append(a["prompt"])

        data = [
            scene["scene_num"], start_t, end_t, dur,
            len(scene["frame_paths"]),
            "", "",  # F、G列留空放图片
            trans, " | ".join(descs), " | ".join(prompts),
        ]
        for col, val in enumerate(data, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.font, c.border, c.fill = body_font, border, fill
            c.alignment = center if col <= 5 else wrap

        # 嵌入关键帧图片：第1帧到F列，第2帧到G列
        ws.row_dimensions[row].height = ROW_HEIGHT_PTS
        for j, fp in enumerate(scene["frame_paths"][:2]):
            if os.path.exists(fp):
                thumb = _make_thumbnail(fp, output_dir, max_h=IMG_HEIGHT)
                img = XlImage(thumb)
                col = IMG_COL_1 if j == 0 else IMG_COL_2
                cell_ref = f"{get_column_letter(col)}{row}"
                ws.add_image(img, cell_ref)

    widths = {"A": 10, "B": 10, "C": 10, "D": 10, "E": 8,
              "F": 28, "G": 28, "H": 35, "I": 50, "J": 60}
    for col_letter, w in widths.items():
        ws.column_dimensions[col_letter].width = w
    ws.freeze_panes = "A2"

    # ── Sheet 2: 概览 ──
    ws2 = wb.create_sheet("概览")
    summary = [
        ("视频标题", video_info.get("title", "未知")),
        ("视频链接", url),
        ("视频时长", format_timestamp(video_info.get("duration", 0))),
        ("总场景数", len(scenes)),
        ("总关键帧数", sum(len(s["frame_paths"]) for s in scenes)),
        ("有语音场景", sum(
            1 for s in scenes
            if transcription and _get_segments_for_scene(
                transcription.get("segments", []), s["start"], s["end"])
        )),
        ("双帧场景(>3s)", sum(1 for s in scenes if len(s["frame_paths"]) > 1)),
    ]
    label_font = Font(bold=True, name="Arial", size=11, color="1F4E79")
    val_font = Font(name="Arial", size=11)
    for r, (label, value) in enumerate(summary, 1):
        ws2.cell(row=r, column=1, value=label).font = label_font
        ws2.cell(row=r, column=2, value=value).font = val_font
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 55

    wb.save(xlsx_path)
    return xlsx_path
